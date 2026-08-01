"""
Stage 1 -- validator.py

Validates raw uploaded bytes BEFORE anything else touches them:
  1. Size limit
  2. Real file-signature/type check (python-magic), not just the extension
  3. Corruption check (can the relevant library actually open it)
  4. Malware scan (ClamAV via pyclamd) if a daemon is reachable, otherwise
     an honestly-labeled integrity/type validation fallback -- we never
     fake a "clean" scan result.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass
from typing import Optional

from app.config import get_settings
from app.models import FileStatus

logger = logging.getLogger("prism.validator")

_SUPPORTED_MIME_EXTENSIONS = {
    "application/pdf": {".pdf"},
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": {".pptx"},
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {".xlsx"},
    "text/csv": {".csv"},
    "text/plain": {".csv"},  # some CSVs are sniffed as text/plain
}


@dataclass
class ValidationResult:
    status: FileStatus
    mime_type: Optional[str]
    detail: str
    malware_scan_method: str  # "clamav" | "integrity_fallback"


def _detect_mime(data: bytes) -> Optional[str]:
    try:
        import magic  # python-magic

        return magic.from_buffer(data, mime=True)
    except Exception:
        logger.warning("python-magic unavailable or failed; falling back to no mime detection")
        return None


def _check_extension_matches_signature(filename: str, mime_type: Optional[str]) -> bool:
    if mime_type is None:
        return True  # can't disprove it, don't hard-fail on this alone
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    allowed_exts = _SUPPORTED_MIME_EXTENSIONS.get(mime_type)
    if allowed_exts is None:
        return False
    return ext in allowed_exts


def _check_corruption(data: bytes, filename: str, mime_type: Optional[str]) -> bool:
    """Returns True if the file opens cleanly with the library that will
    actually parse it in Stage 2."""
    lower = filename.lower()
    try:
        if lower.endswith(".pdf") or mime_type == "application/pdf":
            import fitz  # PyMuPDF

            doc = fitz.open(stream=data, filetype="pdf")
            ok = doc.page_count > 0
            doc.close()
            return ok
        if lower.endswith(".pptx"):
            import io

            from pptx import Presentation

            Presentation(io.BytesIO(data))
            return True
        if lower.endswith(".xlsx"):
            import io

            from openpyxl import load_workbook

            load_workbook(io.BytesIO(data), read_only=True)
            return True
        if lower.endswith(".csv"):
            data.decode("utf-8", errors="strict")
            return True
    except Exception as exc:
        logger.info("Corruption check failed for %s: %s", filename, exc)
        return False
    # Unknown extension: not corrupted per se, just unsupported -- handled
    # by the caller via detected_type == unknown.
    return True


def _scan_malware(data: bytes) -> tuple[bool, str]:
    """
    Returns (is_clean, method). Tries a real ClamAV daemon first. If it's not
    reachable, falls back to a clearly-labeled integrity/type check ONLY --
    this is NOT presented as a malware scan result to the caller.
    """
    try:
        import pyclamd

        cd = pyclamd.ClamdUnixSocket()
        if cd.ping():
            result = cd.scan_stream(data)
            is_clean = result is None
            return is_clean, "clamav"
    except Exception:
        pass
    # Honest fallback: we did NOT run a malware scan. We only assert that
    # the file is structurally valid for its declared type (checked above).
    return True, "integrity_fallback"


def validate_file(filename: str, data: bytes) -> ValidationResult:
    settings = get_settings()

    if len(data) == 0:
        return ValidationResult(FileStatus.CORRUPTED, None, "Empty file.", "integrity_fallback")

    if len(data) > settings.max_upload_size_bytes:
        return ValidationResult(
            FileStatus.CORRUPTED,
            None,
            f"File exceeds max size of {settings.max_upload_size_mb}MB.",
            "integrity_fallback",
        )

    mime_type = _detect_mime(data)

    if not _check_extension_matches_signature(filename, mime_type):
        return ValidationResult(
            FileStatus.CORRUPTED,
            mime_type,
            "File extension does not match its actual signature/content type.",
            "integrity_fallback",
        )

    if not _check_corruption(data, filename, mime_type):
        return ValidationResult(
            FileStatus.CORRUPTED, mime_type, "File is corrupted or unreadable.", "integrity_fallback"
        )

    is_clean, method = _scan_malware(data)
    if not is_clean:
        return ValidationResult(
            FileStatus.MALWARE_FLAGGED, mime_type, "Malware scan flagged this file.", method
        )

    detail = "Passed signature, corruption, and malware checks." if method == "clamav" else (
        "Passed signature and corruption checks. ClamAV daemon was not reachable, "
        "so no live malware scan was performed -- treat this as integrity/type "
        "validation only, not an antivirus guarantee."
    )
    return ValidationResult(FileStatus.VALIDATED, mime_type, detail, method)
